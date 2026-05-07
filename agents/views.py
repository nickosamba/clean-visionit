from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse
from .models import Agent, Incident
from .forms import AgentForm
# incidents/views.py
from clients.models import Client
from accounts.models import User  # Using the custom User model
# Liste des agents
from django.contrib.auth.decorators import login_required
from secteurs.models import Secteur
from accounts.models import User  # Import the custom User model directly
from django.contrib.auth import get_user_model
from django.contrib import messages
from datetime import datetime
from tournees.models import Tournee, TourneeClient
#export
from django.template.loader import render_to_string
from weasyprint import HTML
from django.conf import settings
import datetime as dt
import os
import openpyxl
from openpyxl.styles import Font
User = get_user_model()
from django.core.paginator import Paginator
from core.decorators import agent_required, comptable_required, gestionnaire_required, admin_required

@gestionnaire_required
@login_required
def agent_list(request):
    search = request.GET.get("search", "")
    secteur_filter = request.GET.get("secteur", "")

    agents = Agent.objects.select_related("user", "secteur") \
                          .filter(user__role="Agent de terrain") \
                          .order_by("user__last_name", "user__first_name")

    if search:
        agents = agents.filter(
            user__first_name__icontains=search
        ) | agents.filter(
            user__last_name__icontains=search
        ) | agents.filter(
            matricule__icontains=search
        )

    if secteur_filter:
        agents = agents.filter(secteur__id=secteur_filter)

    secteurs = Secteur.objects.all()

    paginator = Paginator(agents, 10)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "agents": page_obj,
        "page_obj": page_obj,
        "total_agents": agents.count(),
        "secteurs": secteurs,
        "search": search,
        "secteur_filter": secteur_filter,
    }
    return render(request, "agents/agents_list.html", context)

@gestionnaire_required
@login_required
def agent_detail(request, agent_id):
    agent = get_object_or_404(
        Agent.objects.select_related("user", "secteur"),
        id=agent_id
    )

    tournees = (
        Tournee.objects.filter(agent=agent)
        .prefetch_related("clients__client")
        .order_by("-date")
    )

    # Statistiques optimisées
    from django.db.models import Count, Q
    stats = tournees.aggregate(
        total_clients=Count("clients"),
        clients_servis=Count("clients", filter=Q(clients__statut_service="Servi")),
        incidents=Count("clients", filter=Q(clients__statut_service="Problème signalé"))
    )

    total_tournees = tournees.count()
    total_clients = stats["total_clients"]
    clients_servis = stats["clients_servis"]
    incidents = stats["incidents"]
    progression = int((clients_servis / total_clients) * 100) if total_clients > 0 else 0

    # Clients servis (pour la carte)
    clients_points = []
    for t in tournees:
        for tc in t.clients.filter(statut_service="Servi"):
            if tc.client.gps_lat and tc.client.gps_lon:

                # Sécurisation de la rue
                rue_nom = tc.client.rue.nom if tc.client.rue else "Rue inconnue"

                clients_points.append({
                    "nom": tc.client.nom,
                    "lat": tc.client.gps_lat,
                    "lon": tc.client.gps_lon,
                    "adresse": f"{tc.client.adresse_numero} {rue_nom}",
                })

    # Historique incidents
    incidents_list = []
    for t in tournees:
        for tc in t.clients.filter(statut_service="Problème signalé"):

            # Sécurisation de la rue ici aussi
            rue_nom = tc.client.rue.nom if tc.client.rue else "Rue inconnue"

            incidents_list.append({
                "client": tc.client,
                "adresse": f"{tc.client.adresse_numero} {rue_nom}",
                "tournee": t,
                "tc": tc,
            })

    context = {
        "agent": agent,
        "tournees": tournees,
        "total_tournees": total_tournees,
        "total_clients": total_clients,
        "clients_servis": clients_servis,
        "incidents": incidents,
        "progression": progression,
        "clients_points": clients_points,
        "incidents_list": incidents_list,
    }

    return render(request, "agents/agents_detail.html", context)

@gestionnaire_required
@login_required
def agent_update(request, agent_id):
    agent = get_object_or_404(Agent, id=agent_id)
    user = agent.user  # 🔥 Récupération du User lié
    secteurs = Secteur.objects.all()

    if request.method == "POST":

        # 🔥 Mise à jour du User (nom + prénom)
        user.first_name = request.POST.get("first_name", "").strip()
        user.last_name = request.POST.get("last_name", "").strip()
        user.save()

        # 🔥 Mise à jour de l’Agent
        agent.matricule = request.POST.get("matricule")
        agent.telephone_principal = request.POST.get("telephone_principal")
        agent.telephone_secondaire = request.POST.get("telephone_secondaire")
        agent.adresse = request.POST.get("adresse")
        agent.zone_intervention = request.POST.get("zone_intervention")
        agent.statut = request.POST.get("statut")
        agent.type_contrat = request.POST.get("type_contrat")
        agent.date_embauche = request.POST.get("date_embauche")
        agent.secteur_id = request.POST.get("secteur")

        # Photo
        if "photo_profile" in request.FILES:
            agent.photo_profile = request.FILES["photo_profile"]

        agent.save()

        return redirect("agents:detail_agent", agent_id=agent.id)

    return render(request, "agents/agents_update.html", {
        "agent": agent,
        "secteurs": secteurs,
    })

@gestionnaire_required
@login_required
def agent_create(request):
    secteurs = Secteur.objects.all()

    if request.method == "POST":
        email = request.POST.get("email")

        # Vérifier si l'email existe déjà
        if User.objects.filter(email=email).exists():
            return render(request, "agents/ajouter.html", {
                "secteurs": secteurs,
                "Agent": Agent,
                "error": "Cet email est déjà utilisé par un autre utilisateur. Veuillez en choisir un autre.",
            })

        first_name = request.POST.get("first_name", "").strip()
        last_name = request.POST.get("last_name", "").strip()
        telephone_principal = request.POST.get("telephone_principal", "")
        type_contrat = request.POST.get("type_contrat", "")

        if not first_name or not last_name or not telephone_principal or not type_contrat:
            return render(request, "agents/ajouter.html", {
                "secteurs": secteurs,
                "Agent": Agent,
                "error": "Prénom, Nom, Téléphone principal et Type de contrat sont obligatoires.",
            })

        # ✅ Création d’un nouvel utilisateur pour l’agent
        user = User.objects.create_user(
            email=email,
            username=email,
            first_name=first_name,
            last_name=last_name,
            role="Agent de terrain",
            password="clean123"  # mot de passe par défaut
        )

        # Conversion date_embauche
        date_embauche = None
        date_embauche_str = request.POST.get("date_embauche")
        if date_embauche_str:
            try:
                date_embauche = datetime.strptime(date_embauche_str, "%Y-%m-%d").date()
            except ValueError:
                pass

        # Check if an agent was already created by the signal
        try:
            agent = Agent.objects.get(user=user)
            # Update the existing agent with the form data
            agent.telephone_principal = telephone_principal
            agent.telephone_secondaire = request.POST.get("telephone_secondaire", "")
            agent.adresse = request.POST.get("adresse", "")
            agent.zone_intervention = request.POST.get("zone_intervention", "")
            agent.statut = request.POST.get("statut", "Actif")
            agent.type_contrat = type_contrat
            agent.date_embauche = date_embauche
            agent.secteur_id = request.POST.get("secteur")

            if "photo_profile" in request.FILES:
                agent.photo_profile = request.FILES["photo_profile"]

            agent.save()
        except Agent.DoesNotExist:
            # Agent wasn't created by the signal, so create a new one
            agent = Agent.objects.create(
                user=user,
                telephone_principal=telephone_principal,
                telephone_secondaire=request.POST.get("telephone_secondaire", ""),
                adresse=request.POST.get("adresse", ""),
                zone_intervention=request.POST.get("zone_intervention", ""),
                statut=request.POST.get("statut", "Actif"),
                type_contrat=type_contrat,
                date_embauche=date_embauche,
                secteur_id=request.POST.get("secteur"),
            )

            if "photo_profile" in request.FILES:
                agent.photo_profile = request.FILES["photo_profile"]
                agent.save(update_fields=["photo_profile"])

        messages.success(request, f"Agent {user.first_name} {user.last_name} créé avec succès.")
        return redirect("agents:liste")

    return render(request, "agents/ajouter.html", {
        "secteurs": secteurs,
        "Agent": Agent,
    })

@gestionnaire_required
@login_required
def agent_delete(request, agent_id):
    agent = get_object_or_404(Agent, id=agent_id)

    if request.method == "POST":
        user = agent.user
        agent.delete()
        user.delete()
        return redirect("agents:liste")

    return render(request, "agents/agent_delete_confirm.html", {
        "agent": agent
    })

@gestionnaire_required
@login_required
def export_agents_excel(request):
    # Création du classeur Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agents"

    # En-têtes
    headers = [
        "ID",
        "Nom complet",
        "Email",
        "Téléphone principal",
        "Téléphone secondaire",
        "Adresse",
        "Secteur",
        "Zone d’intervention",
        "Statut",
        "Type de contrat",
        "Date d’embauche",
        "Matricule",
    ]

    # Style des en-têtes
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    # Remplissage des données
    agents = Agent.objects.select_related("user", "secteur").all()

    for row, agent in enumerate(agents, start=2):
        ws.cell(row=row, column=1, value=agent.id)
        ws.cell(row=row, column=2, value=f"{agent.user.first_name} {agent.user.last_name}")
        ws.cell(row=row, column=3, value=agent.user.email)
        ws.cell(row=row, column=4, value=agent.telephone_principal)
        ws.cell(row=row, column=5, value=agent.telephone_secondaire)
        ws.cell(row=row, column=6, value=agent.adresse)
        ws.cell(row=row, column=7, value=agent.secteur.nom if agent.secteur else "")
        ws.cell(row=row, column=8, value=agent.zone_intervention)
        ws.cell(row=row, column=9, value=agent.statut)
        ws.cell(row=row, column=10, value=agent.type_contrat)
        ws.cell(row=row, column=11, value=str(agent.date_embauche))
        ws.cell(row=row, column=12, value=agent.matricule)

    # Préparation de la réponse HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="agents_clean.xlsx"'

    wb.save(response)
    return response

@gestionnaire_required
@login_required
def export_agents_pdf(request):
    agents = Agent.objects.select_related("user", "secteur").all()
    user = request.user

    # Chemin du logo
    logo_path = os.path.join(settings.BASE_DIR, "static/img/logo_clean.png")

    html_string = render_to_string("exports/export_pdf.html", {
        "agents": agents,
        "date": dt.date.today().strftime("%d/%m/%Y"),
        "logo_path": logo_path,
        'user': user,
    })

    pdf_file = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = "attachment; filename=agents_clean.pdf"
    return response


# Liste globale des incidents
@gestionnaire_required
@login_required
def liste_incidents(request):
    incidents = Incident.objects.select_related("agent", "client").all().order_by("-date_signalement")
    return render(request, "agents/list_incidents.html", {"incidents": incidents})

# Supprimer un 
@gestionnaire_required
@login_required
def delete_incident(request, id):
    incident = get_object_or_404(Incident, id=id)
    incident.delete()
    return redirect("agents:liste_incidents")

# Détail d’un incident
@gestionnaire_required
@login_required
def detail_incident(request, id):
    incident = get_object_or_404(Incident, id=id)
    return render(request, "agents/detail_incidents.html", {"incident": incident})


from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
from .models import Agent
from django.utils import timezone

def exporter_pdf_agent(request, agent_id):
    agent = get_object_or_404(Agent, id=agent_id)

    html_string = render_to_string("exports/pdf_agent.html", {
        "agent": agent,
        "logo_path": "/static/images/logo.png",  # adapte selon ton projet
        "date": timezone.now().strftime("%d/%m/%Y"),
        "gestionnaire": request.user.get_full_name() or request.user.username,
    })

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename=\"agent_{agent.matricule}.pdf\"'

    HTML(string=html_string).write_pdf(response)
    return response
