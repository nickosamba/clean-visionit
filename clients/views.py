# Create your views here.
from django.shortcuts import render, get_object_or_404, redirect
from django.db import models
from django.urls import reverse
from .models import Client, ClientOption
from .forms import ClientForm
from secteurs.models import Secteur, Rue
from abonnements.models import Abonnement
from paiements.models import Echeance, Paiement
from tournees.models import TourneeClient
from accounts.models import AuditLog
from django.core.paginator import Paginator
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from core.decorators import gestionnaire_required
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
from io import BytesIO
from django.utils import timezone

@gestionnaire_required
@login_required
def client_list(request):
    # Récupération des paramètres GET
    search = request.GET.get("search", "")
    statut = request.GET.get("statut", "")
    abonnement = request.GET.get("abonnement", "")
    page_number = request.GET.get("page", 1)  # numéro de page

    # Base queryset
    clients = Client.objects.select_related("rue", "abonnement").all()

    # Recherche par nom / prénom / code client / téléphone
    if search:
        clients = clients.filter(
            models.Q(nom__icontains=search) |
            models.Q(prenom__icontains=search) |
            models.Q(code_client__icontains=search) |
            models.Q(telephone_principal__icontains=search)
        )

    # Filtre statut
    if statut:
        clients = clients.filter(statut=statut)

    # Filtre abonnement
    if abonnement:
        clients = clients.filter(abonnement_id=abonnement)

    # Pagination (10 clients par page)
    paginator = Paginator(clients, 10)
    page_obj = paginator.get_page(page_number)

    # Envoi au template
    return render(request, "clients/liste.html", {
        "page_obj": page_obj,
        "statuts": Client.Statut.choices,
        "abonnements": Abonnement.objects.all(),
        "search": search,
        "statut": statut,
        "abonnement": abonnement,
    })

@gestionnaire_required
@login_required
def client_details(request, id):
    client = get_object_or_404(Client, id=id)
    return render(request, "clients/details.html", {"client": client})


# @gestionnaire_required
# @login_required
# def client_create(request):
#     if request.method == "POST":
#         form = ClientForm(request.POST)
#         if form.is_valid():
#             client = form.save()

#             # 🔥 Sauvegarde des options sélectionnées
#             selected_options = form.cleaned_data.get("options", [])
#             for option in selected_options:
#                 ClientOption.objects.create(client=client, option=option)

#             # 🔥 Recalcul du montant après ajout des options
#             client.montant_mensuel = client.calculer_montant_mensuel()
#             client.save(update_fields=["montant_mensuel"])

#             return redirect("clients:liste")
#     else:
#         form = ClientForm()
#     return render(request, "clients/ajouter.html", {"form": form, "mode": "create"})


# @gestionnaire_required
# @login_required
# def client_update(request, id):
#     client = get_object_or_404(Client, id=id)

#     if request.method == "POST":
#         form = ClientForm(request.POST, instance=client)
#         if form.is_valid():
#             client = form.save()

#             # 🔥 Supprimer les anciennes options
#             ClientOption.objects.filter(client=client).delete()

#             # 🔥 Ajouter les nouvelles options sélectionnées
#             selected_options = form.cleaned_data.get("options", [])
#             for option in selected_options:
#                 ClientOption.objects.create(client=client, option=option)

#             # 🔥 Recalculer le montant mensuel après ajout des options
#             client.montant_mensuel = client.calculer_montant_mensuel()
#             client.save(update_fields=["montant_mensuel"])

#             return redirect("clients:details", id=id)
#     else:
#         # Pré-remplir les options déjà liées au client
#         initial_options = client.options.values_list("option_id", flat=True)
#         form = ClientForm(instance=client, initial={"options": initial_options})

#     return render(request, "clients/ajouter.html", {
#         "form": form,
#         "client": client
#     })

@gestionnaire_required
@login_required
def client_create(request):
    if request.method == "POST":
        form = ClientForm(request.POST)
        if form.is_valid():
            client = form.save()  # 🔥 Client + Options + montant mensuel gérés dans le form
            
            # 🔒 AuditLog
            AuditLog.objects.create(
                utilisateur=request.user,
                action="Création",
                entite="Client",
                entite_id=client.id,
                details=f"Création du client {client.nom} {client.prenom}"
            )
            
            return redirect("clients:liste")
    else:
        form = ClientForm()
    return render(request, "clients/ajouter.html", {"form": form, "mode": "create"})


@gestionnaire_required
@login_required
def client_update(request, id):
    client = get_object_or_404(Client, id=id)

    if request.method == "POST":
        form = ClientForm(request.POST, instance=client)
        if form.is_valid():
            client = form.save()  # 🔥 tout est géré dans le form
            
            # 🔒 AuditLog
            AuditLog.objects.create(
                utilisateur=request.user,
                action="Modification",
                entite="Client",
                entite_id=client.id,
                details=f"Modification du client {client.nom} {client.prenom}"
            )
            
            return redirect("clients:details", id=id)
    else:
        # Pré-remplir les options déjà liées au client
        initial_options = client.options.values_list("option_id", flat=True)
        form = ClientForm(instance=client, initial={"options": initial_options})

    return render(request, "clients/ajouter.html", {
        "form": form,
        "client": client
    })

@gestionnaire_required
@login_required
def client_delete(request, id):
    client = get_object_or_404(Client, id=id)

    if request.method == "POST":
        client_id = client.id
        client_info = f"{client.nom} {client.prenom}"
        client.delete()
        
        # 🔒 AuditLog
        AuditLog.objects.create(
            utilisateur=request.user,
            action="Suppression",
            entite="Client",
            entite_id=client_id,
            details=f"Suppression du client {client_info}"
        )
        
        return redirect("clients:liste")

    return render(request, "clients/supprimer.html", {"client": client})

#----------------------------------------------------
# Exportation des clients
#----------------------------------------------------
@gestionnaire_required
@login_required
def export_clients_excel(request):
    # Création du classeur
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"

    # En-têtes (sans date_debut)
    headers = [
        "ID",
        "Nom",
        "Prénom",
        "Email",
        "Téléphone",
        "Rue",
        "Abonnement",
        "Montant mensuel",
        "Date recouvrement",
        "Statut",
    ]

    ws.append(headers)

    # Style des en-têtes
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")  # Vert CLEAN
    header_font = Font(color="FFFFFF", bold=True)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Données
    clients = Client.objects.all()

    for client in clients:
        ws.append([
            client.id,
            client.nom,
            client.prenom,
            client.email,
            client.telephone_principal,
            client.rue.nom if client.rue else "",
            client.abonnement.nom if client.abonnement else "",
            client.montant_mensuel,
            client.date_recouvrement,
            client.statut,
        ])

    # Ajustement automatique des colonnes
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    # Réponse HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="clients.xlsx"'
    wb.save(response)

    return response

@gestionnaire_required
@login_required
def export_clients_pdf(request):
    clients = Client.objects.all()

    html_string = render_to_string("exports/export_clients_pdf.html", {
        "clients": clients,
        "request": request,
        "now": timezone.now(),
    })

    # PDF généré en mémoire (fix Windows)
    pdf_file = BytesIO()
    HTML(string=html_string).write_pdf(target=pdf_file)

    pdf_file.seek(0)

    response = HttpResponse(pdf_file.read(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="clients.pdf"'

    return response



from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
from .models import Client
from django.utils import timezone

@gestionnaire_required
@login_required
def exporter_pdf_client(request, client_id):
    client = get_object_or_404(Client, id=client_id)

    # 🔒 AuditLog
    AuditLog.objects.create(
        utilisateur=request.user,
        action="Exportation PDF",
        entite="Client",
        entite_id=client.id,
        details=f"Exportation de la fiche PDF du client {client.nom} {client.prenom}"
    )

    html_string = render_to_string("exports/pdf_client.html", {
        "client": client,
        "logo_path": "/static/images/logo.png",  # adapte selon ton projet
        "date": timezone.now().strftime("%d/%m/%Y"),
        "gestionnaire": request.user.get_full_name() or request.user.username,
    })

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="client_{client.code_client}.pdf"'

    HTML(string=html_string).write_pdf(response)
    return response

