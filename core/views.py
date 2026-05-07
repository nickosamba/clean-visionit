from django.shortcuts import redirect, render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from .models import Setting
from tournees.models import Tournee, TourneeClient
from agents.models import Agent,Incident
from clients.models import Client 
from secteurs.models import Secteur
from abonnements.models import Abonnement
from paiements.models import Echeance, Paiement, Recu
from datetime import date
from django.utils import timezone
from agents.forms import IncidentForm, AgentPhotoForm
from django.template.loader import render_to_string
from weasyprint import HTML, CSS
from django.contrib.auth import get_user_model
from django.core.paginator import Paginator
from django.db.models import Count, Sum, Q
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from .decorators import agent_required, comptable_required, gestionnaire_required, admin_required
from io import BytesIO
from decimal import Decimal
from django.utils.timezone import now
from django.contrib import messages
from .forms import *
User = get_user_model()
from django.core.paginator import Paginator
from django.shortcuts import render
from paiements.models import Recu

#--------------------------------------------
# Partie semi-admin
#-------------------------------------------
@admin_required
@login_required
def dashboard_admin(request):
    # Abonnements
    abonnements_labels = list(Abonnement.objects.values_list("nom", flat=True))
    valeurs_abonnements = [
        Client.objects.filter(abonnement=a).count()
        for a in Abonnement.objects.all()
    ]

    # Agents par secteur
    secteurs_labels = list(Secteur.objects.values_list("nom", flat=True))
    valeurs_agents = [
        Agent.objects.filter(secteur=s).count()
        for s in Secteur.objects.all()
    ]

    # Incidents du jour
    incidents_today = Incident.objects.filter(date_signalement__date=date.today()).count()

    # Incidents du mois (groupés par jour)
    today = date.today()
    incidents_month_qs = (
        Incident.objects.filter(date_signalement__month=today.month, date_signalement__year=today.year)
        .values("date_signalement__day")
        .annotate(total=Count("id"))
        .order_by("date_signalement__day")
    )
    incidents_month_labels = [f"{i['date_signalement__day']}" for i in incidents_month_qs]
    incidents_month_values = [i["total"] for i in incidents_month_qs]

    context = {
        "total_clients": Client.objects.count(),
        "total_agents": Agent.objects.count(),
        "total_secteurs": Secteur.objects.count(),
        "incidents_today": incidents_today,
        "derniers_clients": Client.objects.order_by("-created_at")[:5],
        "abonnements_labels": abonnements_labels,
        "valeurs_abonnements": valeurs_abonnements,
        "secteurs_labels": secteurs_labels,
        "valeurs_agents": valeurs_agents,
        "incidents_month_labels": incidents_month_labels,
        "incidents_month_values": incidents_month_values,
    }
    return render(request, "dashboard/admin.html", context)
@admin_required
@login_required
def users_list(request):
    query = request.GET.get("q", "")

    users = User.objects.filter(is_superuser=False)

    if query:
        users = users.filter(
            Q(username__icontains=query) |
            Q(email__icontains=query) |
            Q(role__icontains=query)
        )

    users = users.order_by("-date_joined")

    paginator = Paginator(users, 10)  # 10 comptes par page
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    return render(request, "admin/users_list.html", {
        "page_obj": page_obj,
        "query": query,
    })

@admin_required
@login_required
def user_detail(request, user_id):
    user = get_object_or_404(User, id=user_id, is_superuser=False)
    return render(request, "admin/user_detail.html", {"user": user})

@admin_required
@login_required
def user_edit(request, user_id):
    user = get_object_or_404(User, id=user_id, is_superuser=False)

    if request.method == "POST":
        form = UserEditForm(request.POST, instance=user)
        if form.is_valid():
            form.save()
            return redirect("dashboard:user_detail", user_id=user.id)
    else:
        form = UserEditForm(instance=user)

    return render(request, "admin/user_edit.html", {
        "form": form,
        "user": user,
    })

@admin_required
@login_required
def user_toggle_active(request, user_id):
    user = get_object_or_404(User, id=user_id, is_superuser=False)

    user.is_active = not user.is_active
    user.save()

    return redirect("dashboard:user_detail", user_id=user.id)

@admin_required
@login_required
def user_delete(request, user_id):
    user = get_object_or_404(User, id=user_id, is_superuser=False)

    if request.method == "POST":
        user.delete()
        return redirect("dashboard:users_list")

    return render(request, "admin/user_delete_confirm.html", {"user": user})

@admin_required
@login_required
def user_change_password(request, user_id):
    user = get_object_or_404(User, id=user_id, is_superuser=False)

    if request.method == "POST":
        form = UserPasswordForm(request.POST)
        if form.is_valid():
            user.set_password(form.cleaned_data["password1"])
            user.save()
            return redirect("dashboard:user_detail", user_id=user.id)
    else:
        form = UserPasswordForm()

    return render(request, "admin/user_change_password.html", {
        "form": form,
        "user": user,
    })
#----------------------------------------------------------
# Dashboard Gestionnaire
#---------------------------------------------------------
@gestionnaire_required
@login_required
def incidents_view(request):
    query = request.GET.get("q", "")
    statut = request.GET.get("statut", "")
    type_inc = request.GET.get("type", "")

    incidents = Incident.objects.select_related("agent", "client").order_by("-date_signalement")

    if query:
        incidents = incidents.filter(
            Q(agent__first_name__icontains=query) |
            Q(agent__last_name__icontains=query) |
            Q(client__nom__icontains=query) |
            Q(description__icontains=query)
        )
    
    if statut:
        incidents = incidents.filter(statut=statut)
    if type_inc:
        incidents = incidents.filter(type_incident=type_inc)

    paginator = Paginator(incidents, 15)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "page_obj": page_obj,
        "query": query,
        "statut_filter": statut,
        "type_filter": type_inc,
        "statut_choices": Incident.StatutIncident.choices,
        "type_choices": Incident.TYPES,
    }
    return render(request, "dashboard/incidents.html", context)

@gestionnaire_required
@login_required
@require_POST
def traiter_incident(request, incident_id):
    incident = get_object_or_404(Incident, id=incident_id)
    nouveau_statut = request.POST.get("statut")
    notes = request.POST.get("notes")

    if nouveau_statut:
        incident.statut = nouveau_statut
    if notes:
        incident.notes_gestionnaire = notes
    
    incident.save()
    messages.success(request, f"L'incident #{incident.id} a été mis à jour.")
    return redirect("dashboard:incidents")

@gestionnaire_required
@login_required
def dashboard_gestionnaire(request):
    today = date.today()

    # 🔍 Tournées du jour
    tournees = (
        Tournee.objects.filter(date=today)
        .select_related("agent", "secteur")
        .prefetch_related("clients__client")
    )

    # 🔍 Clients du jour (FIABLE)
    clients = (
        TourneeClient.objects.filter(tournee__date=today)
        .select_related("client", "tournee", "tournee__agent", "tournee__secteur")
    )

    # Statistiques clients
    total_clients = clients.count()
    served_clients = clients.filter(statut_service="Servi").count()

    # 🟥 Incidents dans les tournées (opérationnels)
    incidents_tournees = clients.filter(statut_service__icontains="probl").count()

    # 🟥 Incidents déclarés (modèle Incident)
    incidents_declares = Incident.objects.filter(
        date_signalement__date=today
    ).count()

    # 🟧 Total incidents unifiés
    incidents = incidents_tournees + incidents_declares

    # 📌 Liste unifiée des incidents
    incidents_list = list(clients.filter(statut_service__icontains="probl"))
    incidents_list += list(Incident.objects.filter(date_signalement__date=today))

    # 🧭 Agents actifs
    agents_actifs = {tc.tournee.agent for tc in clients}

    # 📊 Stats avancées optimisées
    stats_par_statut = {
        "Planifié": clients.filter(statut_service="Planifié").count(),
        "Servi": served_clients,
        "Absent": clients.filter(statut_service="Absent").count(),
        "Problème": incidents,
    }

    # Agrégation par agent (FIABLE et RAPIDE)
    stats_par_agent_qs = (
        clients.filter(statut_service="Servi")
        .values("tournee__agent__user__first_name", "tournee__agent__user__last_name")
        .annotate(total=Count("id"))
    )
    stats_par_agent = {
        f"{item['tournee__agent__user__first_name']} {item['tournee__agent__user__last_name']}": item["total"]
        for item in stats_par_agent_qs
    }

    # Agrégation par secteur
    stats_par_secteur_qs = (
        clients.filter(statut_service="Servi")
        .values("tournee__secteur__nom")
        .annotate(total=Count("id"))
    )
    stats_par_secteur = {item["tournee__secteur__nom"]: item["total"] for item in stats_par_secteur_qs}

    context = {
        "tournees": tournees,
        "clients": clients,
        "progression": int((served_clients / total_clients) * 100) if total_clients else 0,
        "total_clients": total_clients,
        "served_clients": served_clients,
        "incidents": incidents,
        "incidents_list": incidents_list,
        "agents_actifs": agents_actifs,
        "stats_par_statut": stats_par_statut,
        "stats_par_agent": stats_par_agent,
        "stats_par_secteur": stats_par_secteur,
        "today": today,
    }

    return render(request, "dashboard/gestionnaire.html", context)

@gestionnaire_required
@login_required
def export_tournees_pdf(request):
    today = date.today()

    tournees = (
        Tournee.objects.filter(date=today)
        .select_related("agent", "secteur")
        .prefetch_related("clients__client")
    )

    total_clients = sum(t.clients.count() for t in tournees)
    served_clients = sum(t.clients.filter(statut_service="Servi").count() for t in tournees)
    incidents = sum(t.clients.filter(statut_service="Problème signalé").count() for t in tournees)

    html_string = render_to_string("exports/tournees_pdf_gest.html", {
        "tournees": tournees,
        "today": today,
        "total_clients": total_clients,
        "served_clients": served_clients,
        "incidents": incidents,
    })

    pdf_file = HTML(string=html_string).write_pdf(
        stylesheets=[CSS(string="""
            @page { size: A4; margin: 2cm; }
            body { font-family: 'Inter', sans-serif; }
        """)]
    )

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = f"attachment; filename=tournees_{today}.pdf"
    return response


@gestionnaire_required
@login_required
def export_tournees_excel(request):
    today = date.today()

    tournees = (
        Tournee.objects.filter(date=today)
        .select_related("agent", "secteur")
        .prefetch_related("clients__client")
    )

    # Création du fichier Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Tournées du jour"

    # Style premium
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    # En-têtes
    headers = ["Tournée", "Agent", "Secteur", "Client", "Adresse", "Statut"]
    ws.append(headers)

    for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

    # Remplissage des données
    row = 2
    for t in tournees:
        for tc in t.clients.all():
            ws.append([
                f"Tournée {t.id}",
                f"{t.agent.user.first_name} {t.agent.user.last_name}",
                t.secteur.nom,
                tc.client.nom,
                f"{tc.client.adresse_numero} {tc.client.rue.nom}",
                tc.statut_service,
            ])
            for cell in ws[row]:
                cell.border = thin_border
            row += 1

    # Ajustement automatique des colonnes
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    # Réponse HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f"attachment; filename=tournees_{today}.xlsx"
    wb.save(response)

    return response

@gestionnaire_required
@login_required
def export_agents_excel(request):
    agents = Agent.objects.select_related("user", "secteur")

    wb = Workbook()
    ws = wb.active
    ws.title = "Agents"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    headers = ["Nom", "Matricule", "Téléphone", "Secteur", "Statut"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    row = 2
    for a in agents:
        ws.append([
            f"{a.user.first_name} {a.user.last_name}",
            a.matricule,
            a.telephone_principal,
            a.secteur.nom if a.secteur else "—",
            "Actif" if a.user.is_active else "Inactif",
        ])
        for cell in ws[row]:
            cell.border = thin_border
        row += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value)) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=agents.xlsx"
    wb.save(response)
    return response


#-------------------------------------------------
# Dashbaord Agent
#-------------------------------------------------
@agent_required
@login_required
def dashboard_agent(request):
    # Vérifier que l'utilisateur est un agent
    agent = getattr(request.user, "agent", None)
    if agent is None:
        return render(request, "errors/403.html", {
            "message": "Votre profil agent n'est pas encore créé."
        })

    # Récupérer les tournées du jour
    tournees = (
        Tournee.objects
        .filter(agent=agent, date=date.today())
        .prefetch_related("clients__client")
    )

    # Récupérer tous les TourneeClient du jour
    tournee_clients = TourneeClient.objects.filter(
        tournee__in=tournees
    ).select_related("client")

    # Stats
    total_clients = tournee_clients.count()
    served_clients = tournee_clients.filter(statut_service="Servi").count()
    incidents = tournee_clients.filter(statut_service="Problème signalé").count()

    progression = int((served_clients / total_clients) * 100) if total_clients else 0

    context = {
        "agent": agent,
        "tournees": tournees,
        "clients": tournee_clients,
        "total_clients": total_clients,
        "served_clients": served_clients,
        "incidents": incidents,
        "progression": progression,
        "total_incidents": Incident.objects.filter(agent=request.user).count(),
    }

    return render(request, "dashboard/agent.html", context)


@agent_required
@login_required
def confirmer_passage(request, tc_id):
    tc = get_object_or_404(TourneeClient, id=tc_id, tournee__agent=request.user.agent)
    if request.method == "POST":
        tc.statut_service = TourneeClient.StatutService.SERVI
        tc.heure_passage = timezone.now()
        
        # Capture de la photo et du commentaire
        if "photo" in request.FILES:
            tc.photo_preuve = request.FILES["photo"]
        
        commentaire = request.POST.get("commentaire")
        if commentaire:
            tc.commentaires_agent = commentaire
            
        tc.save()
        messages.success(request, f"Passage chez {tc.client.nom} confirmé avec succès.")
    return redirect("dashboard:agent")


@agent_required
@login_required
def signaler_passage(request, tc_id):
    tc = get_object_or_404(TourneeClient, id=tc_id, tournee__agent=request.user.agent)
    if request.method == "POST":
        tc.statut_service = TourneeClient.StatutService.PROBLEME
        tc.save()
        # tu peux aussi ajouter un message Django ici (messages.success)
    return redirect("dashboard:agent")


@agent_required
@login_required
def agent_profile(request):
    agent = request.user.agent  

    if request.method == "POST": 
        form = AgentPhotoForm(request.POST, request.FILES, instance=agent) 
        if form.is_valid(): 
            form.save()
            return redirect("dashboard:agent_profile") 
    else: 
        form = AgentPhotoForm(instance=agent)

    current_month = timezone.now().month

    # Nombre de clients planifiés ce mois
    total_clients = TourneeClient.objects.filter(
        tournee__agent=agent,
        tournee__date__month=current_month
    ).count()

    # Nombre de clients servis
    clients_servis = TourneeClient.objects.filter(
        tournee__agent=agent,
        tournee__date__month=current_month,
        statut_service=TourneeClient.StatutService.SERVI
    ).count()

    # Incidents
    incidents = TourneeClient.objects.filter(
        tournee__agent=agent,
        tournee__date__month=current_month,
        statut_service=TourneeClient.StatutService.PROBLEME
    ).count()

    # Passages = clients servis
    passages = clients_servis

    # Assiduité = % de clients servis / planifiés
    assiduite = int((clients_servis / total_clients) * 100) if total_clients > 0 else 0

    stats = {
        "passages": passages,
        "assiduite": assiduite,
        "total_clients": total_clients,
        "incidents": incidents,
    }

    context = {
        "agent": agent,
        "stats": stats,
        "form": form,
    }
    return render(request, "dashboard/agent_profile.html", context)


@agent_required
@login_required
def agent_incident(request):
    # récupérer uniquement les clients liés aux tournées de l’agent
    clients_agent = Client.objects.filter(tourneeclient__tournee__agent=request.user.agent).distinct()

    if request.method == "POST":
        form = IncidentForm(request.POST, request.FILES)
        form.fields["client"].queryset = clients_agent   # 👈 filtrage
        if form.is_valid():
            incident = form.save(commit=False)
            incident.agent = request.user
            incident.save()
            return redirect("dashboard:agent")
    else:
        form = IncidentForm()
        form.fields["client"].queryset = clients_agent   # 👈 filtrage

    return render(request, "dashboard/agent_incident.html", {"form": form})

#---------------------------------------------------------
# Dashboard Comptable
#---------------------------------------------------------

@comptable_required
@login_required
def dashboard_comptable(request):
    today = date.today()
    mois_courant = today.strftime("%Y-%m")  # ex: "2026-01"

    # 🔢 Statistiques globales avec une seule requête
    stats_globales = Echeance.objects.aggregate(
        total=Count("id"),
        a_payer=Count("id", filter=Q(statut=Echeance.Statut.A_PAYER)),
        payees=Count("id", filter=Q(statut=Echeance.Statut.PAYE)),
        en_retard=Count("id", filter=Q(statut=Echeance.Statut.EN_RETARD)),
    )

    # Paiements du jour (optimisé avec aggregate)
    paiements_today = Paiement.objects.filter(date_paiement__date=today, statut=Paiement.Statut.VALIDE)
    total_paiements_today = paiements_today.aggregate(total=Sum("montant_paye"))["total"] or 0

    # 📊 Stats par mode de paiement (une seule requête)
    stats_par_mode_qs = Paiement.objects.filter(statut=Paiement.Statut.VALIDE).values("mode_paiement").annotate(total=Count("id"))
    stats_par_mode = {item["mode_paiement"]: item["total"] for item in stats_par_mode_qs}

    # ⚠️ Échéances en retard
    echeances_retard = Echeance.objects.filter(statut=Echeance.Statut.EN_RETARD).select_related("client")

    # 📈 Prévisions du mois courant (optimisé avec aggregate)
    montant_attendu_mois = Echeance.objects.filter(periode=mois_courant).aggregate(total=Sum("montant_du"))["total"] or 0
    montant_paye_mois = Paiement.objects.filter(echeance__periode=mois_courant, statut=Paiement.Statut.VALIDE).aggregate(total=Sum("montant_paye"))["total"] or 0
    progression_mois = int((montant_paye_mois / montant_attendu_mois) * 100) if montant_attendu_mois > 0 else 0

    # Préparer les données pour Chart.js
    modes_labels = [mode[1] for mode in Paiement.Mode.choices]
    modes_values = [stats_par_mode.get(mode[0], 0) for mode in Paiement.Mode.choices]

    context = {
        "today": today,
        "total_echeances": stats_globales["total"],
        "a_payer": stats_globales["a_payer"],
        "payees": stats_globales["payees"],
        "en_retard": stats_globales["en_retard"],
        "paiements_today": paiements_today,
        "total_paiements_today": total_paiements_today,
        "stats_par_mode": stats_par_mode,
        "modes_labels": modes_labels,
        "modes_values": modes_values,
        "echeances_retard": echeances_retard,
        "montant_attendu_mois": montant_attendu_mois,
        "montant_paye_mois": montant_paye_mois,
        "progression_mois": progression_mois,
    }
    return render(request, "dashboard/comptable.html", context)

@comptable_required
@login_required
def export_pdf_mensuel(request):
    today = date.today()
    mois_courant = today.strftime("%Y-%m")

    paiements_mois = Paiement.objects.filter(
        echeance__periode=mois_courant,
        statut=Paiement.Statut.VALIDE
    )

    echeances_retard = Echeance.objects.filter(
        periode=mois_courant,
        statut=Echeance.Statut.EN_RETARD
    )

    # Totaux par mode de paiement
    totaux_par_mode = {}
    for mode in Paiement.Mode.choices:
        totaux_par_mode[mode[0]] = sum(
            p.montant_paye for p in paiements_mois.filter(mode_paiement=mode[0])
        )

    context = {
        "mois_courant": mois_courant,
        "today": today,
        "paiements_mois": paiements_mois,
        "echeances_retard": echeances_retard,
        "totaux_par_mode": totaux_par_mode,
    }

    html_string = render_to_string("exports/comptable_pdf.html", context)
    pdf_file = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="rapport_comptable_mensuel_{mois_courant}.pdf"'
    return response

@comptable_required
@login_required
def export_excel_mensuel(request):
    today = date.today()
    mois_courant = today.strftime("%Y-%m")

    paiements_mois = Paiement.objects.filter(
        echeance__periode=mois_courant,
        statut=Paiement.Statut.VALIDE
    )

    # Totaux par mode de paiement
    totaux_par_mode = {}
    for mode in Paiement.Mode.choices:
        totaux_par_mode[mode[0]] = sum(
            p.montant_paye for p in paiements_mois.filter(mode_paiement=mode[0])
        )

    # Création du workbook
    wb = Workbook()

    # --- Onglet 1 : Paiements Mensuels ---
    ws1 = wb.active
    ws1.title = "Paiements Mensuels"
    headers = ["Client", "Montant payé", "Mode", "Date"]
    ws1.append(headers)

    # Style des en-têtes
    for col in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Données
    for paiement in paiements_mois:
        ws1.append([
            paiement.echeance.client.nom,
            float(paiement.montant_paye),
            paiement.mode_paiement,
            paiement.date_paiement.strftime("%d/%m/%Y")
        ])

    # --- Onglet 2 : Synthèse par Mode ---
    ws2 = wb.create_sheet(title="Synthèse par Mode")
    ws2.append(["Mode de paiement", "Total (FCFA)"])

    # Style des en-têtes
    for col in range(1, 3):
        cell = ws2.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Données
    for mode, total in totaux_par_mode.items():
        ws2.append([mode, float(total)])

    # Signature officielle
    row_signature = len(totaux_par_mode) + 3
    ws2.merge_cells(start_row=row_signature, start_column=1, end_row=row_signature, end_column=2)
    signature_cell = ws2.cell(row=row_signature, column=1)
    signature_cell.value = f"Validé par le service comptable CLEAN — Fait à Brazzaville, le {today.strftime('%d/%m/%Y')}"
    signature_cell.font = Font(italic=True, color="555555")
    signature_cell.alignment = Alignment(horizontal="right")

    # Sauvegarde en mémoire
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    # Réponse HTTP
    response = HttpResponse(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename=rapport_comptable_mensuel_{mois_courant}.xlsx'
    return response

# core/views.py

@comptable_required
@login_required
def echeances_list(request):
    today = now().date()

    # Mise à jour automatique des échéances en retard
    Echeance.objects.filter(
        date_echeance__lt=today,
        statut=Echeance.Statut.A_PAYER
    ).update(statut=Echeance.Statut.EN_RETARD)

    echeances = Echeance.objects.select_related("client").order_by("created_at")

    # Pagination
    paginator = Paginator(echeances, 50)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    stats = echeances.aggregate(
        total_payees=Count("id", filter=Q(statut=Echeance.Statut.PAYE)),
        total_retard=Count("id", filter=Q(statut=Echeance.Statut.EN_RETARD)),
        total_a_payer=Count("id", filter=Q(statut=Echeance.Statut.A_PAYER)),
        total_global=Count("id"),
    )

    context = {
        "today": today,
        "page_obj": page_obj,
        **stats
    }
    return render(request, "paiements/echeances.html", context)

# core/views.py
@comptable_required
@login_required
def payer_echeance(request, echeance_id):
    echeance = get_object_or_404(Echeance, id=echeance_id)

    if request.method == "POST":
        montant = Decimal(request.POST.get("montant_paye"))
        mode = request.POST.get("mode_paiement")

        # ✅ Validation
        if montant <= 0:
            messages.error(request, "Le montant payé doit être positif.")
            return redirect("dashboard:payer_echeance", echeance_id=echeance.id)

        if montant > echeance.reste_a_payer:
            messages.error(request, "Le montant payé dépasse le reste à payer.")
            return redirect("dashboard:payer_echeance", echeance_id=echeance.id)

        paiement = Paiement.objects.create(
            echeance=echeance,
            utilisateur=request.user,
            montant_attendu=echeance.montant_du,
            montant_paye=montant,
            mode_paiement=mode,
            statut=Paiement.Statut.VALIDE,
            date_paiement=timezone.now()
        )

        # ✅ Vérifier la somme des paiements liés
        total_paye = Paiement.objects.filter(echeance=echeance).aggregate(
            Sum("montant_paye")
        )["montant_paye__sum"] or Decimal(0)

        if total_paye >= echeance.montant_du:
            echeance.statut = Echeance.Statut.PAYE
        else:
            echeance.statut = Echeance.Statut.A_PAYER
        echeance.save()

        # ✅ Paiement en retard
        if echeance.date_echeance and timezone.now().date() > echeance.date_echeance:
            paiement.notes = "Paiement effectué en retard"
            paiement.save()

        # ✅ Générer automatiquement un reçu lié
        Recu.objects.create(
            paiement=paiement,
            numero_recu=f"REC-{timezone.now().strftime('%Y%m%d')}-{paiement.id}",
            signature_comptable="Comptable CLEAN",
            fichier_url=f"/media/recus/REC-{paiement.id}.pdf"
        )

        messages.success(request, f"Paiement de {montant} FCFA enregistré avec succès.")
        return redirect("dashboard:paiements")


    return render(request, "paiements/payer.html", {
        "echeance": echeance,
        "modes": Paiement.Mode.choices
    })


# core/views.py
from django.shortcuts import render
from django.db.models import Sum
from django.core.paginator import Paginator
from paiements.models import Paiement
from django.utils import timezone
from django.contrib.auth.decorators import login_required
from core.decorators import comptable_required  # adapte selon ton projet

@comptable_required
@login_required
def paiements_list(request):
    # 🔹 Filtres dynamiques
    mode = request.GET.get("mode")
    periode = request.GET.get("periode")  # ex: "mois", "semaine"

    paiements = Paiement.objects.select_related("echeance", "echeance__client").order_by("-date_paiement")

    if mode:
        paiements = paiements.filter(mode_paiement=mode)

    if periode == "mois":
        paiements = paiements.filter(date_paiement__month=timezone.now().month)
    elif periode == "semaine":
        today = timezone.now().date()
        start_week = today - timezone.timedelta(days=today.weekday())
        paiements = paiements.filter(date_paiement__date__gte=start_week)

    # 🔹 Totaux
    total_global = paiements.count()
    total_especes = paiements.filter(mode_paiement=Paiement.Mode.ESPECES).count()
    total_mobile = paiements.filter(mode_paiement=Paiement.Mode.MOBILE_MONEY).count()
    total_virement = paiements.filter(mode_paiement=Paiement.Mode.VIREMENT).count()

    total_montant = paiements.aggregate(Sum("montant_paye"))["montant_paye__sum"] or 0
    total_montant_especes = paiements.filter(mode_paiement=Paiement.Mode.ESPECES).aggregate(Sum("montant_paye"))["montant_paye__sum"] or 0
    total_montant_mobile = paiements.filter(mode_paiement=Paiement.Mode.MOBILE_MONEY).aggregate(Sum("montant_paye"))["montant_paye__sum"] or 0
    total_montant_virement = paiements.filter(mode_paiement=Paiement.Mode.VIREMENT).aggregate(Sum("montant_paye"))["montant_paye__sum"] or 0

    # 🔹 Pagination
    paginator = Paginator(paiements, 20)  # 20 paiements par page
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "page_obj": page_obj,
        "paiements": page_obj.object_list,
        "total_global": total_global,
        "total_especes": total_especes,
        "total_mobile": total_mobile,
        "total_virement": total_virement,
        "total_montant": total_montant,
        "total_montant_especes": total_montant_especes,
        "total_montant_mobile": total_montant_mobile,
        "total_montant_virement": total_montant_virement,
        "today": timezone.now(),
    }
    return render(request, "paiements/paiements.html", context)





from django.db.models import Q
from django.core.paginator import Paginator
from django.shortcuts import render
from paiements.models import Recu

def recu_list(request):
    recus = Recu.objects.select_related(
        "paiement", "paiement__echeance", "paiement__echeance__client"
    ).order_by("-paiement__date_paiement")

    query = request.GET.get("q")
    if query:
        from django.db.models import Q
        recus = recus.filter(
            Q(paiement__echeance__client__full_name__icontains=query) |
            Q(paiement__echeance__periode__icontains=query)
        ).distinct()

    paginator = Paginator(recus, 20)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    return render(request, "paiements/recus.html", {
        "page_obj": page_obj,
        "query": query,
    })


import tempfile

@comptable_required
@login_required
def recu_pdf(request, paiement_id):
    paiement = get_object_or_404(Paiement, id=paiement_id)
    recu = paiement.recus.first()  # récupère le reçu lié

    # Contexte pour le template
    context = {
        "paiement": paiement,
        "recu": recu,
        "logo_path": "static/images/logo_clean.png",  # adapte le chemin
    }

    # Générer HTML
    html_string = render_to_string("exports/recu_pdf.html", context)

    # Générer PDF avec WeasyPrint
    html = HTML(string=html_string)
    result = html.write_pdf()

    # Réponse HTTP
    response = HttpResponse(result, content_type="application/pdf")
    response['Content-Disposition'] = f'inline; filename=recu_{recu.numero_recu}.pdf'
    return response


@require_POST
@login_required
def toggle_theme(request):
    """
    POST endpoint to toggle user theme between 'light' and 'dark'.
    - Updates the Setting model.
    - Updates session for frontend consistency.
    - Returns JSON with current theme.

    Args:
        request: HttpRequest object containing metadata about the request

    Returns:
        JsonResponse: JSON response with theme and success status
    """
    # Get or create user settings
    setting, created = Setting.objects.get_or_create(user=request.user)

    # Toggle the theme value
    if setting.theme == 'light':
        setting.theme = 'dark'
    else:
        setting.theme = 'light'

    # Save the updated theme
    setting.save()
    # Update session for frontend consistency
    request.session['theme'] = setting.theme
    # Return JSON response with theme and success status
    return JsonResponse({'theme': setting.theme, 'success': True})

from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
from django.utils import timezone
from paiements.models import Paiement

def paiements_pdf(request):
    paiements = Paiement.objects.select_related("echeance", "echeance__client").order_by("-date_paiement")

    # Totaux
    from django.db.models import Sum
    total_montant = paiements.aggregate(Sum("montant_paye"))["montant_paye__sum"] or 0

    html_string = render_to_string("exports/paiements_pdf.html", {
        "paiements": paiements,
        "date": timezone.now().strftime("%d/%m/%Y"),
        "gestionnaire": request.user.get_full_name() or request.user.username,
        "logo_path": "/static/images/logo.png",  # adapte selon ton projet
        "total_montant": total_montant,
    })

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="rapport_paiements.pdf"'

    HTML(string=html_string).write_pdf(response)
    return response



from django.shortcuts import render
from django.db.models import Q
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from core.decorators import comptable_required
from paiements.models import Salaire

@comptable_required
@login_required
def salaires_list(request):
    query = request.GET.get("q", "")
    statut = request.GET.get("statut", "")
    periode = request.GET.get("periode", "")

    salaires = Salaire.objects.all().order_by("-periode")

    # 🔎 Recherche globale
    if query:
        salaires = salaires.filter(
            Q(agent__user__username__icontains=query) |
            Q(agent__matricule__icontains=query)
        )

    # 🔎 Filtres
    if statut:
        salaires = salaires.filter(statut=statut)
    if periode:
        salaires = salaires.filter(periode=periode)

    # 🔹 Pagination
    paginator = Paginator(salaires, 10)  # 10 salaires par page
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    return render(request, "paiements/salaires.html", {
        "page_obj": page_obj,
        "query": query,
        "statut": statut,
        "periode": periode,
    })


from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
from django.utils import timezone
from paiements.models import Salaire

def salaires_pdf(request):
    salaires = Salaire.objects.select_related("agent", "agent__user").order_by("-periode")

    # Totaux
    from django.db.models import Sum
    total_montant = salaires.aggregate(Sum("montant"))["montant__sum"] or 0

    html_string = render_to_string("exports/salaires_pdf.html", {
        "salaires": salaires,
        "date": timezone.now().strftime("%d/%m/%Y"),
        "gestionnaire": request.user.get_full_name() or request.user.username,
        "logo_path": "/static/images/logo.png",  # adapte selon ton projet
        "total_montant": total_montant,
    })

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="rapport_salaires.pdf"'

    HTML(string=html_string).write_pdf(response)
    return response





# from django.core.paginator import Paginator
# from django.db.models import Q
# from django.shortcuts import render
# from django.contrib.auth.decorators import login_required
# from core.decorators import role_required
# from paiements.models import Salaire
# from core.decorators import agent_required, comptable_required

# @login_required
# @comptable_required
# def salaires_list(request):
#     query = request.GET.get("q", "")
#     statut = request.GET.get("statut", "")
#     periode = request.GET.get("periode", "")

#     salaires = Salaire.objects.all().order_by("-periode")

#     # 🔎 Recherche globale
#     if query:
#         salaires = salaires.filter(
#             Q(agent__user__username__icontains=query) |
#             Q(agent__matricule__icontains=query)
#         )

#     # 🔎 Filtres
#     if statut:
#         salaires = salaires.filter(statut=statut)
#     if periode:
#         salaires = salaires.filter(periode=periode)

#     # 🔎 Pagination
#     paginator = Paginator(salaires, 10)
#     page_number = request.GET.get("page")
#     page_obj = paginator.get_page(page_number)

#     return render(request, "paiements/salaires.html", {
#         "page_obj": page_obj,
#         "query": query,
#         "statut": statut,
#         "periode": periode,
#     })



@comptable_required
@login_required
def generer_salaires_mensuels_view(request):
    """
    Génère les fiches de salaires 'Prévu' pour tous les agents actifs
    pour le mois en cours.
    """
    from django.db import transaction
    today = now().date()
    # On fixe la période au 1er du mois en cours pour la cohérence
    periode_debut = today.replace(day=1)
    count = 0

    with transaction.atomic():
        agents_actifs = Agent.objects.filter(statut=Agent.Statut.ACTIF)
        
        for agent in agents_actifs:
            # Vérifier si un salaire existe déjà pour cet agent et ce mois
            if not Salaire.objects.filter(agent=agent, periode=periode_debut).exists():
                Salaire.objects.create(
                    agent=agent,
                    montant=agent.salaire_base,
                    periode=periode_debut,
                    statut="Prévu"
                )
                count += 1

    if count > 0:
        messages.success(request, f"{count} fiches de salaires ont été générées pour le mois de {periode_debut.strftime('%B %Y')}.")
    else:
        messages.info(request, "Toutes les fiches de salaires pour ce mois existent déjà.")
        
    return redirect("dashboard:salaires")


@login_required
@comptable_required
def payer_salaire(request, salaire_id):
    salaire = get_object_or_404(Salaire, id=salaire_id)

    if salaire.statut == "Payé":
        messages.warning(request, "Ce salaire est déjà payé.")
        return redirect("dashboard:salaires")

    # ✅ Mise à jour du statut et date
    salaire.statut = "Payé"
    salaire.date_versement = timezone.now().date()
    salaire.save()

    # ✅ Génération d’un reçu (optionnel)
    Recu.objects.create(
        paiement=None, 
        fichier_url=f"/media/recus/SAL-{salaire.id}.pdf",
        numero_recu=f"SAL-{salaire.id}",
        signature_comptable=request.user.get_full_name() or "Service Comptable"
    )

    messages.success(request, f"Le salaire de {salaire.agent} a été validé avec succès.")
    return redirect("dashboard:salaires")


@login_required
@role_required("Comptable")
def ajouter_salaire(request):
    if request.method == "POST":
        form = SalaireForm(request.POST)
        if form.is_valid():
            salaire = form.save(commit=False)
            # ✅ si statut = Payé, on met la date du jour
            if salaire.statut == "Payé" and not salaire.date_versement:
                from django.utils import timezone
                salaire.date_versement = timezone.now().date()
            salaire.save()
            return redirect("dashboard:salaires")
    else:
        form = SalaireForm()
    return render(request, "paiements/ajouter_salaire.html", {"form": form})
